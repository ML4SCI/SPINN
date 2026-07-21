#!/usr/bin/env python3
"""Build a Keynote-friendly PowerPoint deck from hyperparameter sweep outputs."""

from __future__ import annotations

import argparse
import csv
from collections import OrderedDict
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_SUMMARY = PROJECT_ROOT / "SPINN" / "Hyperparameter" / "results" / "summary"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(31, 35, 40)
MUTED = RGBColor(88, 96, 105)
RULE = RGBColor(214, 220, 226)
BLUE = RGBColor(42, 95, 168)
TEAL = RGBColor(0, 132, 130)
GREEN = RGBColor(46, 125, 50)
AMBER = RGBColor(184, 112, 0)
RED = RGBColor(178, 55, 55)
LIGHT_BLUE = RGBColor(232, 241, 252)
LIGHT_GREEN = RGBColor(232, 246, 238)
LIGHT_AMBER = RGBColor(255, 246, 225)
LIGHT_RED = RGBColor(255, 235, 235)
WHITE = RGBColor(255, 255, 255)


def read_rows(summary_csv):
    if not summary_csv.exists():
        return []
    with open(summary_csv, newline="") as f:
        return list(csv.DictReader(f))


def as_float(row, key, default=float("nan")):
    try:
        return float(row.get(key, default))
    except Exception:
        return default


def number(value, digits=3):
    try:
        return f"{float(value):.{digits}g}"
    except Exception:
        return ""


def fit_text(text, max_chars=130):
    text = " ".join(str(text).split())
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip() + "..."


def set_background(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def set_text(run, size=18, color=INK, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def text_box(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(text)
    if align:
        p.alignment = align
    for run in p.runs:
        set_text(run, size=size, color=color, bold=bold)
    return box


def add_title(slide, title, subtitle=None):
    text_box(slide, title, Inches(0.55), Inches(0.28), Inches(9.9), Inches(0.5), size=24, bold=True)
    if subtitle:
        text_box(slide, subtitle, Inches(0.57), Inches(0.78), Inches(10.8), Inches(0.35), size=11, color=MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.06), Inches(12.2), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.color.rgb = RULE


def add_footer(slide, label="SPINN Circular Trap PINN Hyperparameter Study"):
    text_box(slide, label, Inches(0.57), Inches(7.1), Inches(6.2), Inches(0.22), size=8, color=MUTED)


def add_card(slide, x, y, w, h, title, value, fill=LIGHT_BLUE, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = accent
    shape.line.width = Pt(1)
    text_box(slide, title, x + Inches(0.16), y + Inches(0.12), w - Inches(0.32), Inches(0.22), size=8, color=MUTED, bold=True)
    text_box(slide, value, x + Inches(0.16), y + Inches(0.4), w - Inches(0.32), Inches(0.45), size=18, color=accent, bold=True)


def add_bullet_block(slide, title, bullets, x, y, w, h, fill=LIGHT_BLUE, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = accent
    shape.line.width = Pt(1)
    text_box(slide, title, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.3), size=11, color=accent, bold=True)
    box = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.5), w - Inches(0.4), h - Inches(0.65))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = str(bullet)
        p.level = 0
        p.space_after = Pt(3)
        for run in p.runs:
            set_text(run, size=12, color=INK)
    return box


def add_picture_contained(slide, image_path, x, y, max_w, max_h):
    if not image_path.exists():
        return None
    try:
        from PIL import Image

        with Image.open(image_path) as im:
            img_w, img_h = im.size
        ratio = img_w / img_h
        box_ratio = max_w / max_h
        if ratio >= box_ratio:
            w = max_w
            h = max_w / ratio
        else:
            h = max_h
            w = max_h * ratio
        return slide.shapes.add_picture(str(image_path), x + (max_w - w) / 2, y + (max_h - h) / 2, width=w, height=h)
    except Exception:
        return slide.shapes.add_picture(str(image_path), x, y, width=max_w)


def row_by_variant(rows):
    return {r.get("variant_name", ""): r for r in rows}


def best_row(rows, key="phi_mse"):
    valid = [r for r in rows if r.get(key) not in {"", "nan", "NaN", None}]
    return min(valid, key=lambda r: as_float(r, key)) if valid else None


def worst_row(rows, key="phi_mse"):
    valid = [r for r in rows if r.get(key) not in {"", "nan", "NaN", None}]
    return max(valid, key=lambda r: as_float(r, key)) if valid else None


def insight_for_overview(rows):
    best = best_row(rows, "phi_mse")
    base = row_by_variant(rows).get("baseline")
    relu = row_by_variant(rows).get("activation_relu")
    rms = row_by_variant(rows).get("optimizer_rmsprop")
    bullets = []
    if best:
        bullets.append(f"Best DEVSIM agreement came from {best['variant_name']}: Phi MSE {number(best['phi_mse'])}, Psi MSE {number(best['psi_mse'])}.")
    if best and base:
        improvement = as_float(base, "phi_mse") / as_float(best, "phi_mse")
        bullets.append(f"That is about {number(improvement)}x lower RF-potential MSE than the baseline.")
    if relu:
        bullets.append("ReLU is the cautionary case: near-zero training loss, but the worst DEVSIM field error.")
    if rms:
        bullets.append(f"Among optimizer-only variants, RMSprop had the lowest DEVSIM error: Phi MSE {number(rms['phi_mse'])}.")
    return bullets


def graph_insights(rows, graph_name):
    by_name = row_by_variant(rows)
    if graph_name == "normalized_loss_curves":
        relu = by_name.get("activation_relu")
        bc1000 = by_name.get("bc_weight_1000")
        return [
            "Normalized loss lets runs with different starting losses be compared on the same scale.",
            "Fast normalized convergence did not always mean physical accuracy.",
            f"ReLU ended at {number(relu['final_normalized_loss'])} normalized loss but failed the DEVSIM comparison." if relu else "",
            f"BC weight 1000 also converged strongly and was physically accurate: Phi MSE {number(bc1000['phi_mse'])}." if bc1000 else "",
        ]
    if graph_name == "phi_mse":
        best = best_row(rows, "phi_mse")
        worst = worst_row(rows, "phi_mse")
        return [
            f"Lowest RF-potential error: {best['variant_name']} at Phi MSE {number(best['phi_mse'])}." if best else "",
            f"Largest RF-potential error: {worst['variant_name']} at Phi MSE {number(worst['phi_mse'])}." if worst else "",
            "Boundary-condition weighting dominated field accuracy in this first sweep.",
            "Wider networks and sinusoidal activations were the strongest architecture-side options.",
        ]
    if graph_name == "psi_mse":
        best = best_row(rows, "psi_mse")
        return [
            f"Lowest pseudopotential error: {best['variant_name']} at Psi MSE {number(best['psi_mse'])}." if best else "",
            "Pseudopotential depends on gradients, so it punishes field-shape errors more than voltage MSE alone.",
            "The same variants that improved Phi generally improved Psi, but high-frequency instability can still appear in gradients.",
        ]
    if graph_name == "final_normalized_loss":
        return [
            "Final training loss alone is not a reliable model-selection criterion here.",
            "ReLU is artificially favored because the second derivative of ReLU is zero almost everywhere.",
            "The best physical run, bc_weight_1000, also had a low final normalized loss, so it is the most credible winner.",
        ]
    if graph_name == "loss_vs_phi":
        return [
            "Ideal runs sit toward the lower-left: low training loss and low DEVSIM error.",
            "ReLU sits far from the physical winner despite its tiny training loss, exposing a residual-definition failure.",
            "BC weight 1000 is the cleanest point because it pairs low loss with the lowest field error.",
            "Use this graph as the model-selection warning slide: optimization metrics need ground-truth validation.",
        ]
    return []


def sweep_takeaway(sweep_name, sweep_rows):
    ranked = sorted(sweep_rows, key=lambda r: as_float(r, "phi_mse"))
    best = ranked[0]
    if sweep_name == "baseline":
        return [
            f"Baseline gives a useful reference: Phi MSE {number(best['phi_mse'])}, Psi MSE {number(best['psi_mse'])}.",
            "Several variants beat the baseline decisively, so the PINN is hyperparameter-sensitive.",
        ]
    if sweep_name == "width":
        return [
            "Increasing width helped the physical field much more than it helped final training loss.",
            f"Width 256 was best: Phi MSE {number(best['phi_mse'])}, compared with {number(ranked[-1]['phi_mse'])} for the worst width run.",
        ]
    if sweep_name == "depth":
        return [
            "Depth variants were unstable in this setup; deeper did not mean better.",
            f"The best depth-only run was {best['variant_name']}, but it still lagged the baseline by DEVSIM error.",
        ]
    if sweep_name == "activation":
        return [
            "Sinusoidal activation was the best activation-only variant for field accuracy.",
            "ReLU produced misleadingly tiny loss because the Laplace residual becomes degenerate under second derivatives.",
        ]
    if sweep_name == "bc_weight":
        return [
            "Boundary-condition weight was the strongest lever in this sweep.",
            "Too little BC weight let boundary mismatch dominate; 1000 gave the cleanest DEVSIM match.",
        ]
    if sweep_name == "optimizer":
        return [
            "Optimizer choice changed both convergence and physical accuracy.",
            "RMSprop had the best optimizer-only DEVSIM error; SGD plateaued with large boundary error.",
        ]
    if sweep_name == "learning_rate":
        return [
            "Lower learning rates were smoother but too slow for this fixed training budget.",
            "The high-rate run moved fastest but had large PDE spikes; it still beat the lower-rate runs by DEVSIM error.",
        ]
    return [f"Best run in this group: {best['variant_name']}."]


def variant_takeaway(row):
    name = row.get("variant_name", "")
    phi = number(row.get("phi_mse", ""))
    psi = number(row.get("psi_mse", ""))
    norm = number(row.get("final_normalized_loss", ""))
    if name == "activation_relu":
        return "Misleading loss case: the training objective collapses, but the field error is the worst in the sweep."
    if name == "bc_weight_1000":
        return "Best overall DEVSIM match; boundary enforcement was strong without destroying the learned field."
    if name == "width_256":
        return "Best architecture-width result; wider representation improved field shape despite a higher final loss."
    if name == "activation_sinusoidal":
        return "Best activation-only result; smooth derivatives help the PDE residual and gradient-derived pseudopotential."
    if name == "optimizer_sgd":
        return "SGD plateaued; comparison confirms poor field accuracy despite completing the same training budget."
    if name == "baseline":
        return "Reference run for interpreting all variants; several choices beat it by large margins."
    return f"Final norm. loss {norm}; Phi MSE {phi}; Psi MSE {psi}. Use error maps to locate where residual field error concentrates."


def add_title_slide(prs, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    text_box(slide, "SPINN Circular Trap", Inches(0.65), Inches(0.7), Inches(12.0), Inches(0.6), size=34, color=BLUE, bold=True)
    text_box(slide, "PINN hyperparameter study vs. DEVSIM ground truth", Inches(0.68), Inches(1.38), Inches(11.5), Inches(0.5), size=21, color=INK)
    text_box(slide, "Keynote-friendly discussion deck with graph-by-graph interpretation", Inches(0.7), Inches(1.9), Inches(11.0), Inches(0.35), size=13, color=MUTED)

    best = best_row(rows, "phi_mse")
    base = row_by_variant(rows).get("baseline")
    add_card(slide, Inches(0.75), Inches(3.0), Inches(2.75), Inches(1.0), "RUNS COMPLETED", str(len(rows)), LIGHT_BLUE, BLUE)
    add_card(slide, Inches(3.75), Inches(3.0), Inches(2.75), Inches(1.0), "BEST VARIANT", best["variant_name"] if best else "", LIGHT_GREEN, GREEN)
    add_card(slide, Inches(6.75), Inches(3.0), Inches(2.75), Inches(1.0), "BEST PHI MSE", number(best["phi_mse"]) if best else "", LIGHT_GREEN, GREEN)
    if best and base:
        add_card(slide, Inches(9.75), Inches(3.0), Inches(2.75), Inches(1.0), "VS BASELINE", f"{number(as_float(base, 'phi_mse') / as_float(best, 'phi_mse'))}x", LIGHT_AMBER, AMBER)

    add_bullet_block(slide, "Main Readout", insight_for_overview(rows), Inches(0.75), Inches(4.45), Inches(11.8), Inches(1.55), LIGHT_BLUE, BLUE)
    add_footer(slide)


def add_setup_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Study Setup", "Circular Trap DEVSIM ground truth and Vanilla PINN sweep")
    bullets = [
        "Ground truth: Circular Trap DEVSIM RF solve with East/West electrodes at 300 V.",
        "PINN target: solve Laplace equation for Phi_RF; compute RF pseudopotential from autograd gradients.",
        "Comparison: evaluate trained PINN at DEVSIM mesh nodes and compute common-point MSE/error maps.",
        "Sweeps: width, depth, activation, boundary-condition weight, optimizer, and learning rate.",
    ]
    add_bullet_block(slide, "What Was Run", bullets, Inches(0.75), Inches(1.45), Inches(5.75), Inches(4.7), LIGHT_BLUE, BLUE)
    add_bullet_block(
        slide,
        "How To Read The Deck",
        [
            "Convergence plots explain training behavior.",
            "Phi/Psi MSE plots explain agreement with DEVSIM.",
            "Appendix slides show raw PINN minus DEVSIM maps for each variant.",
        ],
        Inches(6.85),
        Inches(1.45),
        Inches(5.75),
        Inches(4.7),
        LIGHT_GREEN,
        GREEN,
    )
    add_footer(slide)


def add_picture_analysis_slide(prs, title, image_path, insights, subtitle=""):
    if not image_path.exists():
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_picture_contained(slide, image_path, Inches(0.45), Inches(1.3), Inches(8.2), Inches(5.35))
    add_bullet_block(slide, "Graph Analysis", [b for b in insights if b], Inches(8.85), Inches(1.45), Inches(3.95), Inches(4.95), LIGHT_AMBER, AMBER)
    add_footer(slide)


def add_top_table_slide(prs, rows):
    ranked = sorted([r for r in rows if r.get("phi_mse")], key=lambda r: as_float(r, "phi_mse"))[:8]
    if not ranked:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Best Runs By DEVSIM Error", "Ranked by RF-potential MSE at DEVSIM mesh points")

    x0, y0 = Inches(0.7), Inches(1.45)
    col_w = [Inches(2.65), Inches(1.35), Inches(1.35), Inches(1.05), Inches(1.05), Inches(1.25), Inches(1.35), Inches(1.35)]
    headers = ["Variant", "Act.", "Opt.", "Width", "Depth", "BC wt.", "Phi MSE", "Psi MSE"]
    row_h = Inches(0.42)
    x = x0
    for j, header in enumerate(headers):
        text_box(slide, header, x, y0, col_w[j], row_h, size=10, color=BLUE, bold=True)
        x += col_w[j]

    for i, row in enumerate(ranked, start=1):
        y = y0 + row_h * i
        fill = LIGHT_GREEN if i == 1 else RGBColor(248, 250, 252)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 - Inches(0.08), y - Inches(0.04), Inches(12.1), row_h)
        bg.fill.solid()
        bg.fill.fore_color.rgb = fill
        bg.line.color.rgb = WHITE
        values = [
            row.get("variant_name", ""),
            row.get("activation", ""),
            row.get("optimizer", ""),
            row.get("width", ""),
            row.get("depth", ""),
            number(row.get("bc_weight", "")),
            number(row.get("phi_mse", "")),
            number(row.get("psi_mse", "")),
        ]
        x = x0
        for j, value in enumerate(values):
            text_box(slide, value, x, y, col_w[j], row_h, size=10, color=INK)
            x += col_w[j]

    add_bullet_block(
        slide,
        "Interpretation",
        [
            "The best run is not just a lower-loss run; it also wins on both Phi and gradient-derived Psi.",
            "The top group includes one BC-weight setting, one width setting, one activation setting, and two optimizer settings.",
        ],
        Inches(0.72),
        Inches(5.75),
        Inches(11.85),
        Inches(0.9),
        LIGHT_BLUE,
        BLUE,
    )
    add_footer(slide)


def add_sweep_summary_slide(prs, sweep_name, sweep_rows):
    ranked = sorted(sweep_rows, key=lambda r: as_float(r, "phi_mse"))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    display_name = "BC Weight" if sweep_name == "bc_weight" else sweep_name.replace("_", " ").title()
    add_title(slide, f"{display_name} Sweep", "Ranked by DEVSIM RF-potential MSE")

    best = ranked[0]
    worst = ranked[-1]
    add_card(slide, Inches(0.7), Inches(1.35), Inches(2.35), Inches(0.95), "BEST", best["variant_name"], LIGHT_GREEN, GREEN)
    add_card(slide, Inches(3.25), Inches(1.35), Inches(2.35), Inches(0.95), "BEST PHI MSE", number(best["phi_mse"]), LIGHT_GREEN, GREEN)
    add_card(slide, Inches(5.8), Inches(1.35), Inches(2.35), Inches(0.95), "BEST PSI MSE", number(best["psi_mse"]), LIGHT_GREEN, GREEN)
    add_card(slide, Inches(8.35), Inches(1.35), Inches(2.35), Inches(0.95), "WORST PHI MSE", number(worst["phi_mse"]), LIGHT_RED, RED)

    headers = ["Variant", "Final Norm. Loss", "AUC Norm. Loss", "Phi MSE", "Psi MSE"]
    col_w = [Inches(3.2), Inches(2.15), Inches(2.15), Inches(1.75), Inches(1.75)]
    x0, y0, row_h = Inches(0.75), Inches(2.75), Inches(0.44)
    x = x0
    for j, header in enumerate(headers):
        text_box(slide, header, x, y0, col_w[j], row_h, size=10, color=BLUE, bold=True)
        x += col_w[j]
    for i, row in enumerate(ranked, start=1):
        y = y0 + row_h * i
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 - Inches(0.08), y - Inches(0.04), Inches(11.1), row_h)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_GREEN if i == 1 else RGBColor(248, 250, 252)
        bg.line.color.rgb = WHITE
        values = [
            row.get("variant_name", ""),
            number(row.get("final_normalized_loss", "")),
            number(row.get("auc_normalized_loss", "")),
            number(row.get("phi_mse", "")),
            number(row.get("psi_mse", "")),
        ]
        x = x0
        for j, value in enumerate(values):
            text_box(slide, value, x, y, col_w[j], row_h, size=10, color=INK)
            x += col_w[j]

    add_bullet_block(slide, "Analysis", sweep_takeaway(sweep_name, ranked), Inches(0.72), Inches(5.55), Inches(11.8), Inches(1.05), LIGHT_AMBER, AMBER)
    add_footer(slide)


def add_discussion_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Discussion Points", "Questions to bring to the group")
    add_bullet_block(
        slide,
        "Interpretation",
        [
            "Do we trust BC-weight 1000 as physically better, or should we test robustness across seeds?",
            "Should ReLU be excluded for second-derivative PINNs, or retained only as a negative control?",
            "Is RMSprop genuinely better here, or should it be followed by LBFGS refinement like many PINN workflows?",
        ],
        Inches(0.75),
        Inches(1.45),
        Inches(5.85),
        Inches(4.85),
        LIGHT_BLUE,
        BLUE,
    )
    add_bullet_block(
        slide,
        "Next Sweep",
        [
            "Repeat top variants with multiple random seeds.",
            "Add Adam to LBFGS and RMSprop to LBFGS two-stage optimizer runs.",
            "Try adaptive loss balancing between PDE and boundary losses.",
        ],
        Inches(6.95),
        Inches(1.45),
        Inches(5.65),
        Inches(4.85),
        LIGHT_GREEN,
        GREEN,
    )
    add_footer(slide)


def add_section_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    text_box(slide, title, Inches(0.75), Inches(2.45), Inches(11.8), Inches(0.75), size=32, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, subtitle, Inches(1.15), Inches(3.28), Inches(11.0), Inches(0.55), size=17, color=INK, align=PP_ALIGN.CENTER)
    add_footer(slide)


def add_variant_error_slides(prs, rows):
    ordered = sorted(rows, key=lambda r: (r.get("sweep_name", ""), r.get("variant_name", "")))
    for row in ordered:
        run_dir = Path(row.get("run_dir", ""))
        phi_plot = run_dir / "comparison_phi_rf.png"
        psi_plot = run_dir / "comparison_psi_rf_eV.png"
        if not phi_plot.exists() and not psi_plot.exists():
            continue

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_background(slide)
        add_title(slide, f"{row.get('variant_name', 'variant')}", "Raw error maps: PINN minus DEVSIM")
        add_picture_contained(slide, phi_plot, Inches(0.35), Inches(1.25), Inches(6.2), Inches(4.75))
        add_picture_contained(slide, psi_plot, Inches(6.65), Inches(1.25), Inches(6.2), Inches(4.75))

        add_card(slide, Inches(0.55), Inches(6.12), Inches(2.25), Inches(0.68), "PHI MSE", number(row.get("phi_mse", "")), LIGHT_GREEN, GREEN)
        add_card(slide, Inches(3.0), Inches(6.12), Inches(2.25), Inches(0.68), "PSI MSE", number(row.get("psi_mse", "")), LIGHT_GREEN, GREEN)
        add_card(slide, Inches(5.45), Inches(6.12), Inches(2.25), Inches(0.68), "NORM. LOSS", number(row.get("final_normalized_loss", "")), LIGHT_BLUE, BLUE)
        text_box(slide, fit_text(variant_takeaway(row), 150), Inches(7.95), Inches(6.16), Inches(4.85), Inches(0.55), size=10, color=INK)
        add_footer(slide)


def write_speaker_guide(rows, out_path):
    by_sweep = OrderedDict()
    for row in rows:
        by_sweep.setdefault(row.get("sweep_name", "unknown"), []).append(row)

    lines = [
        "# SPINN Hyperparameter Results Speaker Guide",
        "",
        "## Main Readout",
    ]
    lines.extend(f"- {bullet}" for bullet in insight_for_overview(rows))
    lines.extend(["", "## Graph Analysis"])
    graph_labels = [
        ("Normalized convergence", "normalized_loss_curves"),
        ("RF potential error", "phi_mse"),
        ("RF pseudopotential error", "psi_mse"),
        ("Final normalized loss", "final_normalized_loss"),
        ("Training loss vs. RF potential error", "loss_vs_phi"),
    ]
    for label, key in graph_labels:
        lines.append(f"### {label}")
        lines.extend(f"- {bullet}" for bullet in graph_insights(rows, key) if bullet)
        lines.append("")
    lines.append("## Sweep Notes")
    for sweep_name, sweep_rows in by_sweep.items():
        lines.append(f"### {sweep_name}")
        lines.extend(f"- {bullet}" for bullet in sweep_takeaway(sweep_name, sweep_rows))
        lines.append("")
    out_path.write_text("\n".join(lines))


def main(argv=None):
    parser = argparse.ArgumentParser(description="Create a SPINN hyperparameter results slide deck.")
    parser.add_argument("--summary_dir", type=Path, default=DEFAULT_SUMMARY)
    parser.add_argument("--out", type=Path, default=DEFAULT_SUMMARY / "spinn_hyperparameter_results.pptx")
    parser.add_argument("--speaker_guide", type=Path, default=DEFAULT_SUMMARY / "spinn_hyperparameter_speaker_guide.md")
    args = parser.parse_args(argv)

    rows = read_rows(args.summary_dir / "summary.csv")
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs, rows)
    add_setup_slide(prs)
    add_picture_analysis_slide(
        prs,
        "Normalized Convergence",
        args.summary_dir / "normalized_loss_curves.png",
        graph_insights(rows, "normalized_loss_curves"),
        "Each curve is total loss divided by that run's initial total loss.",
    )
    add_picture_analysis_slide(
        prs,
        "RF Potential Error",
        args.summary_dir / "phi_mse_by_variant.png",
        graph_insights(rows, "phi_mse"),
        "MSE of Phi_PINN minus Phi_DEVSIM on common DEVSIM mesh nodes.",
    )
    add_picture_analysis_slide(
        prs,
        "RF Pseudopotential Error",
        args.summary_dir / "psi_mse_by_variant.png",
        graph_insights(rows, "psi_mse"),
        "MSE of gradient-derived RF pseudopotential in eV.",
    )
    add_picture_analysis_slide(
        prs,
        "Final Normalized Loss",
        args.summary_dir / "final_normalized_loss_by_variant.png",
        graph_insights(rows, "final_normalized_loss"),
        "Final training loss normalized by each run's initial total loss.",
    )
    add_picture_analysis_slide(
        prs,
        "Training Loss vs. Field Error",
        args.summary_dir / "loss_vs_phi_mse.png",
        graph_insights(rows, "loss_vs_phi"),
        "Log-log comparison of final normalized loss against DEVSIM RF-potential MSE.",
    )
    add_top_table_slide(prs, rows)

    by_sweep = OrderedDict()
    for row in rows:
        by_sweep.setdefault(row.get("sweep_name", "unknown"), []).append(row)
    for sweep_name, sweep_rows in by_sweep.items():
        add_sweep_summary_slide(prs, sweep_name, sweep_rows)

    add_discussion_slide(prs)
    add_section_slide(
        prs,
        "Error Map Appendix",
        "Each slide shows raw PINN minus DEVSIM error maps and a one-line interpretation.",
    )
    add_variant_error_slides(prs, rows)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    write_speaker_guide(rows, args.speaker_guide)
    print(f"Saved slide deck to {args.out}")
    print(f"Saved speaker guide to {args.speaker_guide}")


if __name__ == "__main__":
    main()
